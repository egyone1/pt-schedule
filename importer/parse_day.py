"""Read one day's tab from the clinic booking workbook and turn it into JSON.

Layout of each day tab (as used by customer service):
  row 1  section labels (ot / Inpatient / Aqua / Outpatient / Lym ...)
  row 2  sub-labels or counts
  row 3  therapist names (one column each)
  row 4+ 15-minute slots; time columns hold the clock time,
         a booking is a (usually merged) block of cells in a therapist column.

Phone numbers are dropped on purpose - the website never stores them.
"""
import datetime as dt
import re
import openpyxl

SLOT_MIN = 15
TWO_SESSION_OVER_MIN = 75          # a session longer than 1h15 counts as two

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}

PHONE_RE = re.compile(r"^[\s+()\-⁦⁩\d]{7,}$")
NOTE_WORDS = ("عدم", "ممنوع", "الموافق", "من بعد", "كاش", "دكتور")
JUNK = "\"'$*`⁦⁩"
# blocks that reserve a therapist's time but are not a patient
NOT_PATIENT = {"sports", "sport", "inpatient", "outpatient", "orientation", "meeting", "break", "lunch",
               "off", "training", "lecture", "icu", "ot", "aqua", "lym", "vacation", "اجازة", "اجازه",
               "بريك", "اجتماع", "محاضرة", "محاضره"}


def find_sheet(wb, day):
    """Find the tab whose title is e.g. 'Mon 19 Oct', 'Tues  20 Oct ', 'wed 14 oct'."""
    for ws in wb.worksheets:
        words = ws.title.lower().split()
        nums = [w for w in words if w.isdigit()]
        mons = [MONTHS[w[:3]] for w in words if w[:3] in MONTHS]
        if nums and mons and int(nums[0]) == day.day and mons[0] == day.month:
            return ws
    return None


def clean_name(raw):
    s = re.sub(r"\d+", " ", str(raw))
    s = re.sub(r"[\"'.\n_]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s.title()


def clean_lines(text):
    out = []
    for line in str(text).splitlines():
        line = line.strip().strip(JUNK).strip()
        line = re.sub(r"[\"$`⁦⁩]+", "", line).strip()
        if not line or PHONE_RE.match(line):
            continue
        line = re.sub(r"(\+?\d[\d \-]{8,}\d)", "", line).strip()   # phone inside a line
        if line:
            out.append(line)
    return out


def split_booking(text):
    """-> (patient display name, details) without phone numbers."""
    lines = clean_lines(text)
    names = [l for l in lines if re.search(r"[؀-ۿ]", l) and not l.startswith(NOTE_WORDS)
             and len(l) > 2]
    if not names:                       # Latin-only names such as "Johnson"
        names = [l for l in lines if not l.startswith(NOTE_WORDS)][:1]
    name = names[0] if names else (lines[0] if lines else "")
    # group sessions (aqua) can hold two patients in one block
    extra = [n for n in names[1:] if len(n.split()) >= 2 and not re.search(r"[a-z]", n.lower())
             and not re.match(r"^د\s*[/.\s]", n)]          # "د/احمد" = a doctor, not a patient
    if extra:
        name = name + " + " + " + ".join(extra)
    details = [l for l in lines if l not in names]
    return name, " · ".join(details)


def time_of(v):
    if isinstance(v, dt.time):
        return v
    if isinstance(v, dt.datetime):
        return v.time()
    return None


def parse_day(path_or_wb, day):
    wb = path_or_wb if hasattr(path_or_wb, "worksheets") else openpyxl.load_workbook(path_or_wb)
    ws = find_sheet(wb, day)
    if ws is None:
        return None

    max_row = ws.max_row
    max_col = ws.max_column

    # time columns: most of their cells are clock times
    time_cols = set()
    for c in range(1, max_col + 1):
        n = sum(1 for r in range(4, min(max_row, 60) + 1) if time_of(ws.cell(r, c).value))
        if n >= 8:
            time_cols.add(c)

    # row -> slot time (first time column that has a value on that row)
    row_time = {}
    for r in range(4, max_row + 1):
        for c in sorted(time_cols):
            t = time_of(ws.cell(r, c).value)
            if t:
                row_time[r] = t
                break
    if not row_time:
        return None
    last_row = max(row_time)

    # section label for each column, from row 1
    sections, cur = {}, ""
    for c in range(1, max_col + 1):
        v = ws.cell(1, c).value
        if isinstance(v, str) and v.strip() and v.strip().lower() not in ("eman", "habiba"):
            cur = v.strip()
            cur = cur.upper() if len(cur) <= 3 else cur[0].upper() + cur[1:]
        sections[c] = cur

    merged = {}
    covered = set()
    for m in ws.merged_cells.ranges:
        if m.min_row >= 4:
            merged[(m.min_row, m.min_col)] = m
            for r in range(m.min_row, m.max_row + 1):
                for c in range(m.min_col, m.max_col + 1):
                    if (r, c) != (m.min_row, m.min_col):
                        covered.add((r, c))

    doctors = []
    for c in range(1, max_col + 1):
        head = ws.cell(3, c).value
        if c in time_cols or not isinstance(head, str) or not clean_name(head):
            continue
        sessions = []
        for r in range(4, last_row + 1):
            if (r, c) in covered:
                continue
            v = ws.cell(r, c).value
            if v is None or time_of(v) or not re.search(r"[A-Za-z؀-ۿ]", str(v)):
                continue
            m = merged.get((r, c))
            rows = (m.max_row - m.min_row + 1) if m else 1
            start = row_time.get(r)
            if start is None:
                continue
            name, details = split_booking(v)
            if not name or name.strip().lower() in NOT_PATIENT:
                continue
            dur = rows * SLOT_MIN
            end_dt = dt.datetime.combine(day, start) + dt.timedelta(minutes=dur)
            sessions.append({
                "id": "r%dc%d" % (r, c),
                "patient": name,
                "details": details,
                "start": start.strftime("%H:%M"),
                "end": end_dt.strftime("%H:%M"),
                "minutes": dur,
                "units": 2 if dur > TWO_SESSION_OVER_MIN else 1,
                "status": 0,
            })
        if sessions:
            doctors.append({
                "key": re.sub(r"[^a-z0-9]+", "_", clean_name(head).lower()).strip("_") + "_c%d" % c,
                "name": clean_name(head),
                "section": sections.get(c, ""),
                "order": c,
                "submitted": False,
                "sessions": sessions,
            })
    return {"date": day.isoformat(), "sheet": ws.title.strip(), "doctors": doctors}


if __name__ == "__main__":
    import json, sys
    d = dt.date.fromisoformat(sys.argv[2])
    res = parse_day(sys.argv[1], d)
    with open(sys.argv[3], "w", encoding="utf-8") as f:
        json.dump(res, f, ensure_ascii=False, indent=1)
    if res:
        print(res["sheet"], len(res["doctors"]), "doctors",
              sum(len(x["sessions"]) for x in res["doctors"]), "sessions")
    else:
        print("no sheet for", d)
